import argparse
import pathlib
import subprocess
import sys

from docx import Document
from docx.shared import Cm
from artifact_contract import has_stage, mark_stage_done
from preprocess_md import preprocess_file


def build_docx(book_name: str, md_dir: pathlib.Path) -> pathlib.Path:
    book_dir = md_dir / book_name
    full_md = book_dir / "full.md"
    cover = book_dir / "cover_restored.png"
    out_docx = book_dir / f"{book_name}.docx"
    tmp_docx = book_dir / "_pandoc_temp.docx"
    prep_md = book_dir / "_preprocessed.md"

    if not full_md.exists():
        raise FileNotFoundError(f"full.md не найден: {full_md}")

    # --- 0. Пре-обработка: HTML→pipe-table, \\tag→удалить ---
    preprocess_file(full_md, prep_md)
    print(f"  preprocess: {full_md.name} -> {prep_md.name}")

    # --- 1. Pandoc: md → docx ---
    cmd = [
        "pandoc", str(prep_md),
        "-f", "markdown+pipe_tables+tex_math_dollars",
        "-o", str(tmp_docx),
        "--resource-path", str(book_dir),
        "--toc",
    ]
    print(f"  pandoc: {prep_md.name} -> {tmp_docx.name}")
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(f"pandoc error:\n{res.stderr}")

    # --- 2. Обложка через python-docx ---
    doc = Document(str(tmp_docx))

    if cover.exists():
        print(f"  cover: {cover.name}")
        first_para = doc.paragraphs[0]
        run = first_para.add_run()
        run.add_picture(str(cover), width=Cm(17))
        first_para.alignment = 1  # CENTER
        run.add_break()
    else:
        print(f"  cover: нет (пропуск)")

    # --- 3. Сохранение ---
    doc.save(str(out_docx))
    tmp_docx.unlink(missing_ok=True)
    prep_md.unlink(missing_ok=True)

    mark_stage_done(book_dir, "docx")
    print(f"  -> {out_docx.name}")
    return out_docx


def main():
    parser = argparse.ArgumentParser(description="Сборка docx из full.md + cover_restored.png")
    parser.add_argument("book", nargs="?", help="Имя папки книги (без аргумента — все)")
    args = parser.parse_args()

    md_dir = pathlib.Path("3-md_output")

    if args.book:
        books = [args.book]
    else:
        books = sorted(d.name for d in md_dir.iterdir() if d.is_dir() and (d / "full.md").exists())

    ok, fail = 0, 0
    for name in books:
        print(f"\n[{name}]")
        try:
            book_dir = md_dir / name
            if has_stage(book_dir, "docx"):
                print("  SKIP (уже есть)")
                continue
            build_docx(name, md_dir)
            ok += 1
        except Exception as e:
            print(f"  ERROR: {e}")
            fail += 1

    print(f"\nИтого: {ok} готово, {fail} ошибок")


if __name__ == "__main__":
    main()
